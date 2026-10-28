"""Servicios de fichaje, consolidación diaria e importación de marcaciones."""
import csv
import io
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from zipfile import BadZipFile

from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.empleados.models import Empleado
from apps.horarios_laborales.models import (
    ConfiguracionAsistencia,
    HorarioDiario,
    HorarioPredeterminado,
)
from apps.licencias.models import SolicitudLicencia

from .models import (
    AsignacionJornada,
    Feriado,
    Jornada,
    Marcacion,
    RegistroAsistencia,
)

LIMITE_ARCHIVO = 5 * 1024 * 1024
LIMITE_FILAS = 20_000


def _dias_de_jornada(jornada):
    return {
        int(dia)
        for dia in jornada.dias_laborables.split(',')
        if dia.strip().isdigit()
    }


def _jornada_legacy(empleado, fecha):
    return (
        Jornada.objects.filter(
            empleado=empleado,
            dia_semana=fecha.weekday(),
            vigencia_desde__lte=fecha,
            activa=True,
            turno__activo=True,
        )
        .filter(Q(vigencia_hasta__isnull=True) | Q(vigencia_hasta__gte=fecha))
        .select_related('turno')
        .order_by('-vigencia_desde')
        .first()
    )


def jornada_de(empleado):
    """Devuelve la jornada asignada, la predeterminada o la semanal vigente."""
    asignacion = (
        AsignacionJornada.objects.filter(empleado=empleado, jornada__activa=True)
        .select_related('jornada')
        .first()
    )
    if asignacion:
        return asignacion.jornada
    predeterminada = Jornada.objects.filter(
        empleado__isnull=True, es_predeterminada=True, activa=True,
    ).first()
    if predeterminada:
        return predeterminada
    return _jornada_legacy(empleado, timezone.localdate())


def _jornada_para_fecha(empleado, fecha):
    try:
        asignacion = empleado.asignacion_jornada
    except AsignacionJornada.DoesNotExist:
        asignacion = None
    if asignacion and asignacion.jornada.activa:
        return asignacion.jornada
    predeterminada = getattr(empleado, '_jornada_predeterminada', None)
    if predeterminada and predeterminada.activa:
        return predeterminada
    jornadas = getattr(empleado, '_jornadas_legacy', None)
    if jornadas is not None:
        for jornada in jornadas:
            if (
                jornada.empleado_id
                and jornada.activa
                and jornada.dia_semana == fecha.weekday()
                and jornada.vigencia_desde <= fecha
                and (jornada.vigencia_hasta is None or jornada.vigencia_hasta >= fecha)
                and jornada.turno_id
                and jornada.turno.activo
            ):
                return jornada
        return None
    return _jornada_legacy(empleado, fecha)


@dataclass(frozen=True)
class HorarioEsperado:
    inicio: datetime | None
    fin: datetime | None
    es_franco: bool
    fuente: str
    tolerancia_minutos: int


_SIN_HORARIO_DIARIO_EN_CACHE = object()
_SIN_HORARIO_PREDETERMINADO_EN_CACHE = object()
_SIN_VALOR_EN_CACHE = object()
_TIPOS_AUTOMATICOS = {
    RegistroAsistencia.Tipo.PRESENTE,
    RegistroAsistencia.Tipo.AUSENTE,
    RegistroAsistencia.Tipo.TARDANZA,
    RegistroAsistencia.Tipo.RETIRO_ANTICIPADO,
}


def _combinar_fecha_hora(fecha, hora):
    resultado = datetime.combine(fecha, hora)
    if settings.USE_TZ:
        return timezone.make_aware(resultado, timezone.get_current_timezone())
    return resultado


def _configuracion_asistencia():
    return ConfiguracionAsistencia.objects.filter(pk=1).first() or ConfiguracionAsistencia()


def _construir_intervalo(fecha, hora_entrada, hora_salida):
    inicio = _combinar_fecha_hora(fecha, hora_entrada)
    fin = _combinar_fecha_hora(fecha, hora_salida)
    if fin <= inicio:
        fin += timedelta(days=1)
    return inicio, fin


def obtener_horario_esperado(empleado, fecha):
    """Resuelve primero el horario diario y después la jornada existente."""
    horario_diario = getattr(
        empleado,
        '_horario_diario',
        _SIN_HORARIO_DIARIO_EN_CACHE,
    )
    if horario_diario is _SIN_HORARIO_DIARIO_EN_CACHE:
        horario_diario = HorarioDiario.objects.filter(
            empleado=empleado,
            fecha=fecha,
        ).first()

    if horario_diario is not None:
        config = getattr(empleado, '_configuracion_asistencia', None)
        if config is None:
            config = _configuracion_asistencia()
        if horario_diario.es_franco:
            return HorarioEsperado(
                inicio=None,
                fin=None,
                es_franco=True,
                fuente='diario',
                tolerancia_minutos=config.tolerancia_tardanza_minutos,
            )
        inicio, fin = _construir_intervalo(
            fecha,
            horario_diario.hora_entrada,
            horario_diario.hora_salida,
        )
        return HorarioEsperado(
            inicio=inicio,
            fin=fin,
            es_franco=False,
            fuente='diario',
            tolerancia_minutos=config.tolerancia_tardanza_minutos,
        )

    jornada = _jornada_para_fecha(empleado, fecha)
    if jornada is None:
        return None
    horario_predeterminado = getattr(
        empleado,
        '_horario_predeterminado',
        _SIN_HORARIO_PREDETERMINADO_EN_CACHE,
    )
    if horario_predeterminado is _SIN_HORARIO_PREDETERMINADO_EN_CACHE:
        horario_predeterminado = HorarioPredeterminado.objects.filter(
            empleado=empleado,
            vigente_desde__lte=fecha,
        ).order_by('-vigente_desde').first()
    if (
        horario_predeterminado is not None
        and not _es_dia_laborable(empleado, jornada, fecha)
    ):
        return None
    if horario_predeterminado is not None:
        hora_entrada = horario_predeterminado.hora_entrada
        hora_salida = horario_predeterminado.hora_salida
        tolerancia = jornada.tolerancia_minutos
        if jornada.turno_id and jornada.hora_entrada is None:
            tolerancia = jornada.turno.minutos_tolerancia
    elif jornada.hora_entrada is not None and jornada.hora_salida is not None:
        hora_entrada = jornada.hora_entrada
        hora_salida = jornada.hora_salida
        tolerancia = jornada.tolerancia_minutos
    elif jornada.turno_id:
        hora_entrada = jornada.turno.hora_inicio
        hora_salida = jornada.turno.hora_fin
        tolerancia = jornada.turno.minutos_tolerancia
    else:
        return None

    inicio, fin = _construir_intervalo(fecha, hora_entrada, hora_salida)
    return HorarioEsperado(
        inicio=inicio,
        fin=fin,
        es_franco=False,
        fuente='predeterminado' if horario_predeterminado else 'jornada',
        tolerancia_minutos=tolerancia,
    )


def _marca_horaria(fecha_hora):
    local = timezone.localtime(fecha_hora) if settings.USE_TZ else fecha_hora
    return local.time().replace(tzinfo=None)


def _obtener_datetimes_reales(registro, esperado):
    ingreso = (
        _combinar_fecha_hora(registro.fecha, registro.hora_ingreso)
        if registro.hora_ingreso else None
    )
    egreso = (
        _combinar_fecha_hora(registro.fecha, registro.hora_egreso)
        if registro.hora_egreso else None
    )
    if esperado and esperado.inicio and esperado.fin:
        if (
            ingreso
            and esperado.fin.date() > registro.fecha
            and ingreso < esperado.inicio
        ):
            ingreso += timedelta(days=1)
        if (
            egreso
            and ingreso
            and egreso <= ingreso
            and esperado.fin.date() > registro.fecha
        ):
            egreso += timedelta(days=1)
        elif (
            egreso
            and ingreso is None
            and esperado.fin.date() > registro.fecha
            and egreso <= esperado.fin
        ):
            egreso += timedelta(days=1)
    elif egreso and ingreso and egreso <= ingreso:
        egreso += timedelta(days=1)
    return ingreso, egreso


def _hay_licencia_o_vacaciones(registro):
    cache = getattr(registro, '_tiene_licencia_o_vacaciones', None)
    if cache is not None:
        return cache
    return _tipo_licencia_en_fecha(registro) is not None


def _tipo_licencia_en_fecha(registro):
    cache = getattr(registro, '_tipo_licencia_o_vacaciones', _SIN_VALOR_EN_CACHE)
    if cache is not _SIN_VALOR_EN_CACHE:
        return cache
    descuenta_vacaciones = SolicitudLicencia.objects.filter(
        empleado_id=registro.empleado_id,
        fecha_inicio__lte=registro.fecha,
        fecha_fin__gte=registro.fecha,
    ).values_list('tipo_licencia__descuenta_vacaciones', flat=True).first()
    if descuenta_vacaciones is None:
        return None
    return (
        RegistroAsistencia.Tipo.VACACIONES
        if descuenta_vacaciones
        else RegistroAsistencia.Tipo.LICENCIA
    )


def _es_feriado(registro):
    cache = getattr(registro, '_es_feriado', None)
    if cache is not None:
        return cache
    return Feriado.objects.filter(fecha=registro.fecha).exists()


def _agregar_observacion(registro, observacion):
    if observacion not in registro.observaciones.splitlines():
        registro.observaciones = (
            f'{registro.observaciones}\n{observacion}'.strip()
            if registro.observaciones else observacion
        )


def calcular_novedades(registro):
    """Actualiza estado y métricas de forma repetible usando un único cálculo."""
    empleado = registro.empleado
    esperado = getattr(registro, '_horario_esperado', _SIN_VALOR_EN_CACHE)
    if esperado is _SIN_VALOR_EN_CACHE:
        esperado = obtener_horario_esperado(empleado, registro.fecha)

    horario_diario = getattr(registro, '_horario_diario_calculado', _SIN_VALOR_EN_CACHE)
    if horario_diario is _SIN_VALOR_EN_CACHE:
        horario_diario = getattr(
            empleado,
            '_horario_diario',
            _SIN_HORARIO_DIARIO_EN_CACHE,
        )
    if horario_diario is _SIN_HORARIO_DIARIO_EN_CACHE:
        horario_diario = HorarioDiario.objects.filter(
            empleado_id=registro.empleado_id,
            fecha=registro.fecha,
        ).first()
    registro.horario_diario = horario_diario

    es_feriado = _es_feriado(registro) or registro.tipo == RegistroAsistencia.Tipo.FERIADO
    tipo_licencia = _tipo_licencia_en_fecha(registro)
    tiene_licencia = tipo_licencia is not None
    ingreso, egreso = _obtener_datetimes_reales(registro, esperado)
    tiene_fichaje = ingreso is not None or egreso is not None
    minutos_tardanza = 0
    minutos_extra = 0
    es_tarde = False
    retiro_anticipado = False

    config = getattr(registro, '_configuracion_asistencia', None)
    if config is None:
        config = _configuracion_asistencia()

    if esperado and esperado.inicio and ingreso is not None:
        # La tardanza cuenta solo los minutos posteriores al fin de la tolerancia.
        tolerancia = timedelta(minutes=esperado.tolerancia_minutos)
        limite = esperado.inicio + tolerancia
        diferencia = int((ingreso - limite).total_seconds() // 60)
        if diferencia > 0:
            minutos_tardanza = diferencia
            es_tarde = True

    if esperado and esperado.fin and egreso is not None:
        retiro_anticipado = egreso < esperado.fin

    duracion_trabajada = 0
    if ingreso is not None and egreso is not None and egreso > ingreso:
        duracion_trabajada = int((egreso - ingreso).total_seconds() // 60)

    if tiene_fichaje and (es_feriado or (esperado and esperado.es_franco)):
        minutos_extra = duracion_trabajada
    elif esperado and esperado.fin and egreso is not None:
        exceso = int((egreso - esperado.fin).total_seconds() // 60)
        # Política temporal: el mínimo global aplica a horarios diarios; Jornada mantiene
        # el comportamiento histórico de contar cualquier exceso positivo.
        minimo_extra = (
            config.minimo_extra_minutos
            if esperado.fuente == 'diario'
            else 0
        )
        if exceso >= minimo_extra and exceso > 0:
            minutos_extra = exceso

    if not tiene_fichaje:
        if (
            esperado is not None
            and not esperado.es_franco
            and not es_feriado
            and not tiene_licencia
        ):
            registro.estado = RegistroAsistencia.Estado.AUSENTE
        else:
            registro.estado = RegistroAsistencia.Estado.SIN_HORARIO
    elif es_feriado:
        registro.estado = RegistroAsistencia.Estado.FERIADO_TRABAJADO
    elif esperado and esperado.es_franco:
        registro.estado = RegistroAsistencia.Estado.FRANCO_TRABAJADO
    elif esperado is None:
        registro.estado = RegistroAsistencia.Estado.SIN_HORARIO
    else:
        registro.estado = (
            RegistroAsistencia.Estado.TARDE
            if es_tarde
            else RegistroAsistencia.Estado.PUNTUAL
        )

    if (
        registro.tipo in _TIPOS_AUTOMATICOS
        and tiene_fichaje
        and tiene_licencia
        and not es_feriado
        and tipo_licencia is not None
    ):
        registro.tipo = tipo_licencia
    elif registro.tipo in _TIPOS_AUTOMATICOS and tiene_fichaje:
        if es_tarde:
            registro.tipo = RegistroAsistencia.Tipo.TARDANZA
        elif (
            retiro_anticipado
            and not es_feriado
            and esperado is not None
            and not esperado.es_franco
        ):
            registro.tipo = RegistroAsistencia.Tipo.RETIRO_ANTICIPADO
        else:
            registro.tipo = RegistroAsistencia.Tipo.PRESENTE

    if registro.tipo != RegistroAsistencia.Tipo.HOME_OFFICE:
        registro.minutos_tardanza = minutos_tardanza
        registro.minutos_extra = minutos_extra

    if tiene_fichaje and tiene_licencia:
        _agregar_observacion(
            registro,
            'Marcación realizada durante una solicitud de licencia o vacaciones.',
        )

    return registro


def _construir_consolidacion(registro, marcaciones, ahora=None):
    marcaciones = sorted(marcaciones, key=lambda marca: marca.fecha_hora)
    registro.hora_ingreso = (
        _marca_horaria(marcaciones[0].fecha_hora) if marcaciones else None
    )
    registro.hora_egreso = (
        _marca_horaria(marcaciones[-1].fecha_hora)
        if len(marcaciones) > 1
        else None
    )

    if registro.tipo not in {
        RegistroAsistencia.Tipo.FERIADO,
        RegistroAsistencia.Tipo.HOME_OFFICE,
        RegistroAsistencia.Tipo.HORA_EXTRA,
    }:
        if not registro.hora_egreso:
            nota = 'Sin marcación de salida'
            if nota not in registro.observaciones:
                registro.observaciones = (
                    f'{registro.observaciones}\n{nota}'.strip()
                    if registro.observaciones else nota
                )
        else:
            registro.observaciones = '\n'.join(
                linea for linea in registro.observaciones.splitlines()
                if linea != 'Sin marcación de salida'
            )

    if ahora is not None:
        registro.fecha_actualizacion = ahora


def _rango_dia(fecha):
    zona = timezone.get_current_timezone()
    inicio = datetime.combine(fecha, time.min)
    fin = datetime.combine(fecha + timedelta(days=1), time.min)
    if settings.USE_TZ:
        inicio = timezone.make_aware(inicio, zona)
        fin = timezone.make_aware(fin, zona)
    return inicio, fin


def _adjuntar_contexto_calculo(
    registro,
    esperado,
    horario_diario,
    feriado,
    tiene_licencia,
    configuracion,
    tipo_licencia=None,
):
    registro._horario_esperado = esperado
    registro._horario_diario_calculado = horario_diario
    registro._es_feriado = feriado
    registro._tiene_licencia_o_vacaciones = tiene_licencia
    registro._configuracion_asistencia = configuracion
    registro._tipo_licencia_o_vacaciones = tipo_licencia


def consolidar_dia(empleado, fecha):
    """Actualiza el resumen diario usando las marcaciones de hora local."""
    esperado = obtener_horario_esperado(empleado, fecha)
    inicio, fin = _rango_dia(fecha)
    marcaciones = list(
        Marcacion.objects.filter(
            empleado=empleado,
            fecha_hora__gte=inicio,
            fecha_hora__lt=fin,
        ).order_by('fecha_hora')
    )
    if esperado and esperado.fin and esperado.fin.date() > fecha:
        inicio_siguiente, fin_siguiente = _rango_dia(fecha + timedelta(days=1))
        marcaciones.extend(Marcacion.objects.filter(
            empleado=empleado,
            tipo=Marcacion.Tipo.SALIDA,
            fecha_hora__gte=inicio_siguiente,
            fecha_hora__lt=fin_siguiente,
        ).order_by('fecha_hora'))
    if not marcaciones:
        return None
    registro, _ = RegistroAsistencia.objects.get_or_create(
        empleado=empleado,
        fecha=fecha,
        defaults={'tipo': RegistroAsistencia.Tipo.PRESENTE},
    )
    horario_diario = HorarioDiario.objects.filter(
        empleado=empleado,
        fecha=fecha,
    ).first()
    tipo_licencia = _tipo_licencia_en_fecha(registro)
    _adjuntar_contexto_calculo(
        registro,
        esperado,
        horario_diario,
        Feriado.objects.filter(fecha=fecha).exists(),
        tipo_licencia is not None,
        _configuracion_asistencia(),
        tipo_licencia,
    )
    _construir_consolidacion(registro, marcaciones, timezone.now())
    registro.save()
    return registro


def recalcular_asistencias(empleado_id, fecha):
    """Recalcula todos los registros existentes de un empleado para una fecha."""
    registros = RegistroAsistencia.objects.filter(
        empleado_id=empleado_id,
        fecha=fecha,
    ).select_related('empleado')
    cantidad = 0
    for registro in registros:
        registro.save()
        cantidad += 1
    return cantidad


@transaction.atomic
def registrar_marcacion(
    empleado,
    origen,
    ip=None,
    momento=None,
    cargado_por=None,
    tipo=None,
):
    """Registra entrada o salida usando la hora del servidor y consolida el día."""
    momento = momento or timezone.now()
    if timezone.is_naive(momento):
        raise ValidationError('La fecha y hora de la marcación debe incluir zona horaria.')
    momento = momento.astimezone(timezone.get_current_timezone())
    Empleado.objects.select_for_update().get(pk=empleado.pk)
    ultima = Marcacion.objects.select_for_update().filter(
        empleado=empleado,
    ).order_by('-fecha_hora').first()
    if ultima and (momento - ultima.fecha_hora).total_seconds() < 60:
        raise ValidationError('Espera al menos 60 segundos antes de volver a marcar.')

    inicio, fin = _rango_dia(timezone.localtime(momento).date())
    marcas_hoy = Marcacion.objects.filter(
        empleado=empleado,
        fecha_hora__gte=inicio,
        fecha_hora__lt=fin,
    )
    tiene_entrada = marcas_hoy.filter(tipo=Marcacion.Tipo.ENTRADA).exists()
    if tipo is None:
        tipo = Marcacion.Tipo.SALIDA if tiene_entrada else Marcacion.Tipo.ENTRADA
    elif tipo == Marcacion.Tipo.ENTRADA and tiene_entrada:
        raise ValidationError('La llegada ya fue registrada para hoy.')
    elif tipo == Marcacion.Tipo.SALIDA and not tiene_entrada:
        raise ValidationError('No se puede registrar la salida antes de la llegada.')
    elif tipo not in Marcacion.Tipo.values:
        raise ValidationError('El tipo de marcación no es válido.')
    marcacion = Marcacion.objects.create(
        empleado=empleado,
        fecha_hora=momento,
        tipo=tipo,
        origen=origen,
        ip=ip,
        cargado_por=cargado_por,
    )
    consolidar_dia(empleado, timezone.localtime(momento).date())
    return marcacion


def _es_dia_laborable(empleado, jornada, fecha):
    if jornada is None:
        return False
    if jornada.hora_entrada is not None:
        return fecha.weekday() in _dias_de_jornada(jornada)
    return bool(jornada.turno_id)


def _consolidar_en_lote(registro, marcas, jornada, ahora):
    _construir_consolidacion(registro, marcas, ahora)
    calcular_novedades(registro)
    return registro


def cerrar_dia(fecha):
    """Consolida las marcaciones y crea ausencias sin realizar consultas por empleado."""
    from django.db.models import Prefetch

    inicio, fin = _rango_dia(fecha)
    ahora = timezone.now()
    jornada_predeterminada = Jornada.objects.filter(
        empleado__isnull=True, es_predeterminada=True, activa=True,
    ).first()
    empleados = list(
        Empleado.objects.filter(estado=Empleado.Estado.ACTIVO)
        .select_related('asignacion_jornada__jornada')
        .prefetch_related(Prefetch(
            'jornadas',
            queryset=Jornada.objects.filter(empleado__isnull=False)
            .select_related('turno').order_by('-vigencia_desde'),
            to_attr='_jornadas_legacy',
        ))
    )
    for empleado in empleados:
        empleado._jornada_predeterminada = jornada_predeterminada
    empleado_ids = [empleado.pk for empleado in empleados]
    if not empleado_ids:
        return {
            'consolidados': 0,
            'ausentes_creados': 0,
            'licencias_creadas': 0,
            'omitidos': 0,
        }

    configuracion = _configuracion_asistencia()
    horarios_predeterminados = {}
    for horario in HorarioPredeterminado.objects.filter(
        empleado_id__in=empleado_ids,
        vigente_desde__lte=fecha,
    ).order_by('empleado_id', '-vigente_desde'):
        horarios_predeterminados.setdefault(horario.empleado_id, horario)
    horarios_diarios = {
        horario.empleado_id: horario
        for horario in HorarioDiario.objects.filter(
            empleado_id__in=empleado_ids,
            fecha=fecha,
        )
    }
    for empleado in empleados:
        empleado._horario_diario = horarios_diarios.get(empleado.pk)
        empleado._horario_predeterminado = horarios_predeterminados.get(
            empleado.pk,
        )
        empleado._configuracion_asistencia = configuracion

    marcas_por_empleado = defaultdict(list)
    marcas = Marcacion.objects.filter(
        empleado_id__in=empleado_ids,
        fecha_hora__gte=inicio,
        fecha_hora__lt=fin,
    ).order_by('fecha_hora')
    for marca in marcas:
        marcas_por_empleado[marca.empleado_id].append(marca)
    inicio_siguiente, fin_siguiente = _rango_dia(fecha + timedelta(days=1))
    salidas_siguiente_dia = defaultdict(list)
    for marca in Marcacion.objects.filter(
        empleado_id__in=empleado_ids,
        tipo=Marcacion.Tipo.SALIDA,
        fecha_hora__gte=inicio_siguiente,
        fecha_hora__lt=fin_siguiente,
    ).order_by('fecha_hora'):
        salidas_siguiente_dia[marca.empleado_id].append(marca)

    registros = {
        registro.empleado_id: registro
        for registro in RegistroAsistencia.objects.filter(
            empleado_id__in=empleado_ids, fecha=fecha,
        )
    }
    licencias = {
        empleado_id: es_vacaciones
        for empleado_id, es_vacaciones in SolicitudLicencia.objects.filter(
            empleado_id__in=empleado_ids,
            fecha_inicio__lte=fecha,
            fecha_fin__gte=fecha,
        ).values_list(
            'empleado_id',
            'tipo_licencia__descuenta_vacaciones',
        )
    }
    feriado = Feriado.objects.filter(fecha=fecha).exists()

    consolidados_existentes = []
    consolidados_nuevos = []
    registros_actualizados = []
    ausentes = []
    licencias_creadas = []
    omitidos = 0
    for empleado in empleados:
        empleado_marcas = list(marcas_por_empleado.get(empleado.pk, []))
        registro = registros.get(empleado.pk)
        jornada = _jornada_para_fecha(empleado, fecha)
        esperado = obtener_horario_esperado(empleado, fecha)
        if esperado and esperado.fin and esperado.fin.date() > fecha:
            empleado_marcas.extend(salidas_siguiente_dia.get(empleado.pk, []))
        tiene_licencia = empleado.pk in licencias
        tipo_licencia = (
            RegistroAsistencia.Tipo.VACACIONES
            if licencias.get(empleado.pk)
            else RegistroAsistencia.Tipo.LICENCIA
            if tiene_licencia
            else None
        )
        if empleado_marcas:
            if registro is None:
                registro = RegistroAsistencia(
                    empleado=empleado,
                    fecha=fecha,
                    tipo=RegistroAsistencia.Tipo.PRESENTE,
                )
                consolidados_nuevos.append(registro)
            _adjuntar_contexto_calculo(
                registro,
                esperado,
                horarios_diarios.get(empleado.pk),
                feriado,
                tiene_licencia,
                configuracion,
                tipo_licencia,
            )
            _consolidar_en_lote(registro, empleado_marcas, jornada, ahora)
            if registro.pk:
                consolidados_existentes.append(registro)
            continue
        if empleado.pk in licencias:
            es_vacaciones = licencias[empleado.pk]
            if es_vacaciones and (fecha.weekday() >= 5 or feriado):
                omitidos += 1
                continue
            tipo_licencia = (
                RegistroAsistencia.Tipo.VACACIONES
                if es_vacaciones
                else RegistroAsistencia.Tipo.LICENCIA
            )
            if registro is not None:
                if registro.tipo == RegistroAsistencia.Tipo.AUSENTE:
                    registro.tipo = tipo_licencia
                    _adjuntar_contexto_calculo(
                        registro,
                        esperado,
                        horarios_diarios.get(empleado.pk),
                        feriado,
                        True,
                        configuracion,
                        tipo_licencia,
                    )
                    calcular_novedades(registro)
                    registros_actualizados.append(registro)
                else:
                    omitidos += 1
                continue
            nuevo_registro = RegistroAsistencia(
                empleado=empleado,
                fecha=fecha,
                tipo=tipo_licencia,
            )
            _adjuntar_contexto_calculo(
                nuevo_registro,
                esperado,
                horarios_diarios.get(empleado.pk),
                feriado,
                True,
                configuracion,
                tipo_licencia,
            )
            calcular_novedades(nuevo_registro)
            licencias_creadas.append(nuevo_registro)
            continue
        if registro is not None:
            omitidos += 1
            continue
        if feriado:
            omitidos += 1
            continue
        if esperado is not None and esperado.es_franco:
            omitidos += 1
            continue
        if jornada is not None and not _es_dia_laborable(empleado, jornada, fecha):
            omitidos += 1
            continue
        if jornada is None and fecha.weekday() >= 5:
            omitidos += 1
            continue
        nuevo_ausente = RegistroAsistencia(
            empleado=empleado,
            fecha=fecha,
            tipo=RegistroAsistencia.Tipo.AUSENTE,
        )
        _adjuntar_contexto_calculo(
            nuevo_ausente,
            esperado,
            horarios_diarios.get(empleado.pk),
            feriado,
            False,
            configuracion,
            None,
        )
        calcular_novedades(nuevo_ausente)
        ausentes.append(nuevo_ausente)

    if consolidados_existentes:
        RegistroAsistencia.objects.bulk_update(
            consolidados_existentes,
            [
                'hora_ingreso', 'hora_egreso', 'tipo', 'minutos_tardanza',
                'minutos_extra', 'observaciones', 'fecha_actualizacion',
                'estado', 'horario_diario',
            ],
        )
    if registros_actualizados:
        RegistroAsistencia.objects.bulk_update(
            registros_actualizados,
            [
                'tipo', 'observaciones', 'minutos_tardanza', 'minutos_extra',
                'fecha_actualizacion', 'estado', 'horario_diario',
            ],
        )
    if consolidados_nuevos or ausentes or licencias_creadas:
        RegistroAsistencia.objects.bulk_create(
            consolidados_nuevos + ausentes + licencias_creadas,
            ignore_conflicts=True,
        )
    return {
        'consolidados': len(consolidados_existentes) + len(consolidados_nuevos),
        'ausentes_creados': len(ausentes),
        'licencias_creadas': len(licencias_creadas) + len(registros_actualizados),
        'omitidos': omitidos,
    }


def _normalizar_cabecera(value):
    texto = unicodedata.normalize('NFD', str(value or ''))
    texto = ''.join(char for char in texto if unicodedata.category(char) != 'Mn')
    return '_'.join(texto.strip().lower().replace('-', ' ').split())


def _leer_archivo(archivo):
    contenido = archivo.read()
    if len(contenido) > LIMITE_ARCHIVO:
        raise ValidationError('El archivo no puede superar los 5 MB.')
    nombre = getattr(archivo, 'name', '').lower()
    if nombre.endswith('.xlsx'):
        try:
            libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        except (InvalidFileException, BadZipFile, OSError) as exc:
            raise ValidationError('El archivo XLSX no es válido o está dañado.') from exc
        try:
            filas = libro.active.iter_rows(values_only=True)
            cabecera = next(filas, None)
            if cabecera is None:
                raise ValidationError('El archivo está vacío.')
            return list(cabecera), list(filas)
        finally:
            libro.close()
    if not nombre.endswith('.csv'):
        raise ValidationError('El archivo debe ser CSV o XLSX.')
    try:
        texto = contenido.decode('utf-8-sig')
    except UnicodeDecodeError:
        texto = contenido.decode('latin-1')
    muestra = texto[:8192]
    try:
        delimitador = csv.Sniffer().sniff(muestra, delimiters=',;').delimiter
    except csv.Error:
        delimitador = ';' if muestra.count(';') > muestra.count(',') else ','
    lector = csv.reader(io.StringIO(texto), delimiter=delimitador)
    cabecera = next(lector, None)
    if cabecera is None:
        raise ValidationError('El archivo está vacío.')
    return cabecera, list(lector)


def _parsear_hora(valor):
    if isinstance(valor, datetime):
        return valor.time()
    if isinstance(valor, time):
        return valor
    texto = str(valor or '').strip()
    for formato in ('%H:%M:%S', '%H:%M'):
        try:
            return datetime.strptime(texto, formato).time()
        except ValueError:
            pass
    raise ValueError('formato de hora inválido')


def _parsear_fecha_hora(valor_fecha, valor_hora=None):
    if isinstance(valor_fecha, datetime):
        momento = (
            datetime.combine(valor_fecha.date(), _parsear_hora(valor_hora))
            if valor_hora is not None else valor_fecha
        )
    elif isinstance(valor_fecha, date):
        momento = datetime.combine(
            valor_fecha,
            _parsear_hora(valor_hora) if valor_hora is not None else time.min,
        )
    else:
        texto = str(valor_fecha or '').strip()
        momento = None
        for formato in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M',
                        '%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M',
                        '%Y-%m-%dT%H:%M:%S', '%Y-%m-%dT%H:%M'):
            try:
                momento = datetime.strptime(texto, formato)
                break
            except ValueError:
                pass
        if momento is None and ('T' in texto or ' ' in texto):
            try:
                momento = datetime.fromisoformat(texto.replace('Z', '+00:00'))
            except ValueError:
                pass
        if momento is None:
            for formato in ('%d/%m/%Y', '%Y-%m-%d'):
                try:
                    fecha = datetime.strptime(texto, formato).date()
                    break
                except ValueError:
                    fecha = None
            if fecha is None:
                raise ValueError('formato de fecha inválido')
            hora = _parsear_hora(valor_hora)
            momento = datetime.combine(fecha, hora)
    if timezone.is_naive(momento):
        momento = timezone.make_aware(momento, timezone.get_current_timezone())
    return momento.astimezone(timezone.get_current_timezone())


def _tipo_importacion(valor):
    valor = _normalizar_cabecera(valor).upper()
    tipos = {
        'E': Marcacion.Tipo.ENTRADA,
        'ENTRADA': Marcacion.Tipo.ENTRADA,
        'S': Marcacion.Tipo.SALIDA,
        'SALIDA': Marcacion.Tipo.SALIDA,
    }
    if valor not in tipos:
        raise ValueError('tipo inválido; usa ENTRADA, SALIDA, E o S')
    return tipos[valor]


@transaction.atomic
def importar_marcaciones(archivo, usuario):
    """Importa marcaciones CSV/XLSX, reporta errores de fila y consolida sus días."""
    cabecera, filas = _leer_archivo(archivo)
    if len(filas) > LIMITE_FILAS:
        raise ValidationError('El archivo no puede contener más de 20.000 filas.')
    indices = {_normalizar_cabecera(nombre): indice for indice, nombre in enumerate(cabecera)}
    if 'legajo' not in indices:
        raise ValidationError('El archivo debe incluir una columna legajo.')
    if 'fecha_hora' not in indices and not {'fecha', 'hora'}.issubset(indices):
        raise ValidationError('Incluye fecha_hora o ambas columnas fecha y hora.')

    filas_validas = []
    errores = []
    legajos = set()
    for numero, fila in enumerate(filas, start=2):
        def valor(nombre):
            indice = indices.get(nombre)
            return fila[indice] if indice is not None and indice < len(fila) else None

        legajo = str(valor('legajo') or '').strip()
        try:
            if not legajo:
                raise ValueError('falta el legajo')
            fecha_hora = (
                _parsear_fecha_hora(valor('fecha_hora'))
                if 'fecha_hora' in indices
                else _parsear_fecha_hora(valor('fecha'), valor('hora'))
            )
            tipo = _tipo_importacion(valor('tipo')) if str(valor('tipo') or '').strip() else None
            filas_validas.append((numero, legajo, fecha_hora, tipo))
            legajos.add(legajo)
        except (ValueError, TypeError, OverflowError) as exc:
            errores.append({'fila': numero, 'motivo': str(exc)})

    empleados_por_legajo = {
        empleado.legajo: empleado
        for empleado in Empleado.objects.filter(legajo__in=legajos)
    }
    candidatos = []
    por_dia = defaultdict(list)
    for numero, legajo, fecha_hora, tipo in filas_validas:
        empleado = empleados_por_legajo.get(legajo)
        if empleado is None:
            errores.append({'fila': numero, 'motivo': f'no existe el legajo {legajo}'})
            continue
        item = (numero, empleado, fecha_hora, tipo)
        if tipo is None:
            por_dia[(empleado.pk, timezone.localtime(fecha_hora).date())].append(item)
        else:
            candidatos.append(item)

    ignoradas = 0
    for grupo in por_dia.values():
        grupo.sort(key=lambda item: item[2])
        seleccion = [grupo[0]]
        if len(grupo) > 1:
            seleccion.append(grupo[-1])
            ignoradas += max(0, len(grupo) - 2)
        for indice, item in enumerate(seleccion):
            numero, empleado, fecha_hora, _ = item
            tipo = Marcacion.Tipo.ENTRADA if indice == 0 else Marcacion.Tipo.SALIDA
            candidatos.append((numero, empleado, fecha_hora, tipo))

    candidatos.sort(key=lambda item: (item[1].pk, item[2]))
    in_file = set()
    unicos = []
    duplicadas = 0
    for numero, empleado, fecha_hora, tipo in candidatos:
        identidad = (empleado.pk, fecha_hora)
        if identidad in in_file:
            duplicadas += 1
            continue
        in_file.add(identidad)
        unicos.append((numero, empleado, fecha_hora, tipo))

    if unicos:
        empleados_ids = {empleado.pk for _, empleado, _, _ in unicos}
        minimo = min(fecha_hora for _, _, fecha_hora, _ in unicos)
        maximo = max(fecha_hora for _, _, fecha_hora, _ in unicos)
        existentes = set(
            Marcacion.objects.filter(
                empleado_id__in=empleados_ids,
                fecha_hora__range=(minimo, maximo),
            ).values_list('empleado_id', 'fecha_hora')
        )
        nuevas = []
        for _, empleado, fecha_hora, tipo in unicos:
            identidad = (empleado.pk, fecha_hora)
            if identidad in existentes:
                duplicadas += 1
                continue
            nuevas.append(Marcacion(
                empleado=empleado,
                fecha_hora=fecha_hora,
                tipo=tipo,
                origen=Marcacion.Origen.IMPORTACION,
                cargado_por=usuario,
            ))
        Marcacion.objects.bulk_create(nuevas, ignore_conflicts=True)
        dias_afectados = {
            (empleado, timezone.localtime(fecha_hora).date())
            for _, empleado, fecha_hora, _ in unicos
        }
        for empleado, fecha in dias_afectados:
            consolidar_dia(empleado, fecha)
    else:
        nuevas = []
        dias_afectados = set()

    return {
        'creadas': len(nuevas),
        'duplicadas': duplicadas,
        'ignoradas': ignoradas,
        'errores': sorted(errores, key=lambda error: error['fila']),
        'dias_consolidados': len(dias_afectados),
    }
