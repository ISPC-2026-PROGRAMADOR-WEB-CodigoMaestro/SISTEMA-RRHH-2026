"""Utilidades para generar reportes en PDF (reportlab) y Excel (openpyxl)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

AZUL = colors.HexColor('#2c3e50')
GRIS = colors.HexColor('#f2f2f2')


def generar_pdf(titulo, columnas, filas, subtitulo=''):
    """Devuelve un buffer BytesIO con el reporte en PDF."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        title=titulo, leftMargin=24, rightMargin=24,
        topMargin=24, bottomMargin=24,
    )
    estilos = getSampleStyleSheet()
    elementos = [Paragraph(titulo, estilos['Title'])]
    if subtitulo:
        elementos.append(Paragraph(subtitulo, estilos['Normal']))
    elementos.append(Spacer(1, 12))

    data = [columnas] + [
        ['' if c is None else str(c) for c in fila] for fila in filas
    ]
    tabla = Table(data, repeatRows=1)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2c3e50')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f2f2f2')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elementos.append(tabla)
    doc.build(elementos)
    buffer.seek(0)
    return buffer


def generar_excel(titulo, columnas, filas):
    """Devuelve un buffer BytesIO con el reporte en Excel (.xlsx)."""
    wb = Workbook()
    ws = wb.active
    ws.title = (titulo or 'Reporte')[:31]
    ws.append([titulo])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columnas))
    ws['A1'].font = Font(bold=True, size=13)
    ws['A1'].alignment = Alignment(horizontal='center')

    ws.append(columnas)
    for celda in ws[2]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='2C3E50')
        celda.alignment = Alignment(horizontal='center')

    for fila in filas:
        ws.append(['' if v is None else str(v) for v in fila])

    for i, col in enumerate(ws.iter_cols(min_row=2), start=1):
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = min(max_len + 3, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
